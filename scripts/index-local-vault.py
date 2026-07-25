"""Build a local, derived text/metadata index for ADAMA source vaults.

The source vault is immutable. Outputs are intentionally gitignored.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


MAX_TEXT_PER_FILE = 30_000
MAX_PDF_PAGES = 25
MAX_PARSE_BYTES = 40 * 1024 * 1024
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".srt", ".vtt"}
THEMES = {
    "adama_brand": ("adama", "адама", "cba", "сва", "logo", "логотип"),
    "jewish_calendar": (
        "shabbat", "шаббат", "pesach", "песах", "sukkot", "суккот",
        "hanukkah", "chanukah", "ханук", "purim", "пурим", "rosh hashana",
        "рош а-шана", "ту би шват", "tu b", "lag baomer", "лаг баомер",
    ),
    "hebrew_learning": ("hebrew", "иврит", "עברית", "ulpan", "ульпан"),
    "aliyah_integration": ("aliyah", "алия", "репат", "new immigrant", "олим", "עולים"),
    "community_events": (
        "event", "мероприят", "встреч", "поездк", "trip", "shabbaton",
        "шаббатон", "charidy", "чариди", "community", "общин",
    ),
    "social_content": (
        "instagram", "reels", "story", "stories", "пост", "сторис",
        "реклама", "promo", "publication", "פרסומים",
    ),
    "rebbe_chabad": ("rebbe", "ребе", "любавич", "chabad", "хабад", "חב״ד"),
}


def normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()[:MAX_TEXT_PER_FILE]


def extract(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    try:
        if path.stat().st_size > MAX_PARSE_BYTES:
            return "", "metadata:size-limit"
        if ext == ".pdf":
            reader = PdfReader(str(path))
            return normalize(
                "\n".join((page.extract_text() or "") for page in reader.pages[:MAX_PDF_PAGES])
            ), "pdf"
        if ext == ".docx":
            doc = Document(str(path))
            return normalize("\n".join(p.text for p in doc.paragraphs)), "docx"
        if ext == ".pptx":
            deck = Presentation(str(path))
            chunks = []
            for slide in deck.slides:
                chunks.extend(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            return normalize("\n".join(chunks)), "pptx"
        if ext in {".xlsx", ".xlsm"}:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            chunks = []
            for sheet in workbook.worksheets:
                chunks.append(f"[Sheet: {sheet.title}]")
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if row_number > 500:
                        break
                    chunks.append(" | ".join("" if value is None else str(value) for value in row))
            workbook.close()
            return normalize("\n".join(chunks)), "xlsx"
        if ext in TEXT_EXTENSIONS:
            return normalize(path.read_text(encoding="utf-8", errors="replace")), "text"
        return "", "metadata"
    except Exception as exc:  # A single damaged file must not abort a vault scan.
        return "", f"error:{type(exc).__name__}"


def main() -> int:
    if len(sys.argv) != 3:
        print("usage: index-local-vault.py MANIFEST.csv OUTPUT_DIR", file=sys.stderr)
        return 2

    manifest_path = Path(sys.argv[1]).resolve()
    output_dir = Path(sys.argv[2]).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    format_counts: Counter[str] = Counter()
    method_counts: Counter[str] = Counter()
    theme_counts: Counter[str] = Counter()
    theme_samples: dict[str, list[str]] = defaultdict(list)
    source_counts: Counter[str] = Counter()
    errors: list[dict[str, str]] = []
    records = 0
    extracted_files = 0
    extracted_characters = 0

    ndjson_path = output_dir / "extracted.ndjson"
    with manifest_path.open("r", encoding="utf-8-sig", newline="") as source, ndjson_path.open(
        "w", encoding="utf-8"
    ) as target:
        reader = csv.DictReader(source)
        for row in reader:
            records += 1
            source_name = row["source"]
            relative_path = row["relative_path"]
            # Manifest paths already include the source root (Downloads/ or work/).
            vault_path = manifest_path.parent / "vault-imports" / Path(relative_path)
            ext = vault_path.suffix.lower() or "[no extension]"
            format_counts[ext] += 1
            source_counts[source_name] += 1

            text, method = extract(vault_path)
            method_counts[method] += 1
            if method.startswith("error:"):
                errors.append({"path": f"{source_name}/{relative_path}", "error": method})
            if text:
                extracted_files += 1
                extracted_characters += len(text)

            searchable = f"{relative_path}\n{text}".casefold()
            matched_themes = []
            for theme, terms in THEMES.items():
                if any(term.casefold() in searchable for term in terms):
                    theme_counts[theme] += 1
                    matched_themes.append(theme)
                    if len(theme_samples[theme]) < 12:
                        theme_samples[theme].append(f"{source_name}/{relative_path}")

            target.write(
                json.dumps(
                    {
                        "source": source_name,
                        "relative_path": relative_path,
                        "extension": ext,
                        "bytes": int(row["bytes"]),
                        "modified": row["modified"],
                        "extraction_method": method,
                        "themes": matched_themes,
                        "text": text,
                    },
                    ensure_ascii=False,
                )
                + "\n"
            )

    summary = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "manifest": str(manifest_path),
        "records": records,
        "source_counts": dict(source_counts.most_common()),
        "format_counts": dict(format_counts.most_common()),
        "method_counts": dict(method_counts.most_common()),
        "extracted_files": extracted_files,
        "extracted_characters": extracted_characters,
        "theme_counts": dict(theme_counts.most_common()),
        "theme_samples": dict(theme_samples),
        "errors_count": len(errors),
        "errors": errors[:100],
        "limitations": [
            "Images, audio, video, archives, installers, and legacy binary Office files are metadata-indexed only.",
            "PDF extraction is text-layer based; scanned PDFs require a later OCR pass.",
            f"PDF extraction is limited to the first {MAX_PDF_PAGES} pages.",
            f"Files larger than {MAX_PARSE_BYTES // (1024 * 1024)} MiB are metadata-indexed only.",
            f"Extracted text is capped at {MAX_TEXT_PER_FILE} characters per file.",
        ],
    }
    (output_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f"Indexed {records} files; extracted text from {extracted_files}; "
        f"errors: {len(errors)}; summary: {output_dir / 'summary.json'}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
